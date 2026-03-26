"""DLP Gateway — FastAPI v3 with Document Classification, DPDP, JWT Auth"""
from __future__ import annotations

import hashlib
import io
import logging
import os
import uuid
from datetime import datetime
from typing import List, Optional

from fastapi import (
    FastAPI, WebSocket, WebSocketDisconnect,
    Depends, HTTPException, Query, UploadFile, File, Form,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import OAuth2PasswordRequestForm
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession

from db.session import get_session, engine
from db.models import Base, DLPEvent, UserRiskProfile, Alert
from detection.engine import dlp_engine
from detection.classifier import doc_classifier, RESTRICTED, CONFIDENTIAL, INTERNAL, PUBLIC
from policy.engine import policy_engine
from cache.redis_client import get_cached, set_cached
from alerting.alert_manager import alert_manager
from gateway.auth import get_current_user, require_admin, create_access_token, USERS_DB, verify_password

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("dlp.gateway")

app = FastAPI(title="LLM DLP Gateway", version="3.0.0", docs_url="/api/docs")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

UI_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "ui")
if os.path.isdir(UI_DIR):
    app.mount("/ui", StaticFiles(directory=UI_DIR, html=True), name="ui")


# ── Startup ───────────────────────────────────────────────────────────────────
@app.on_event("startup")
async def startup():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    logger.info("DLP Gateway v3 started — Document Classification + JWT enabled")


# ── Auth ──────────────────────────────────────────────────────────────────────
@app.post("/auth/login")
async def login(form_data: OAuth2PasswordRequestForm = Depends()):
    user = USERS_DB.get(form_data.username)
    if not user or not verify_password(form_data.password, user["hashed_password"]):
        raise HTTPException(status_code=401, detail="Incorrect username or password")
    token = create_access_token({
        "sub": form_data.username,
        "role": user["role"],
        "department": user["department"],
    })
    return {"access_token": token, "token_type": "bearer", "role": user["role"]}


# ── Schemas ───────────────────────────────────────────────────────────────────
class AnalyzeRequest(BaseModel):
    user_id:           str   = Field(..., example="emp-00123")
    department:        str   = Field("default", example="finance")
    role:              str   = Field("employee", example="analyst")
    prompt:            str   = Field(..., min_length=1)
    destination_model: str   = Field("unknown", example="chatgpt")


class AnalyzeResponse(BaseModel):
    event_id:           str
    decision:           str
    risk_score:         float
    risk_tier:          str
    block_reason:       str
    detected_types:     List[str]
    findings:           List[dict]
    layer_scores:       dict
    processing_ms:      float
    policy_notes:       dict
    from_cache:         bool
    timestamp:          str
    dpdp_violation:     bool = False
    dpdp_categories:    List[str] = []
    llm_triggered:      bool = False


# ── DPDP Compliance Map ───────────────────────────────────────────────────────
DPDP_SENSITIVE_MAP = {
    "Aadhaar Number":      "biometric_id",
    "PAN Number":          "financial_id",
    "Email Address":       "contact_data",
    "Date of Birth":       "sensitive_personal",
    "person name":         "personal_identifier",
    "home address":        "contact_data",
    "phone number":        "contact_data",
    "Salary Figure":       "financial_personal",
    "bank account number": "financial_personal",
    "Passport Number":     "government_id",
    "US SSN":              "government_id",
    "UK NI Number":        "government_id",
    "India Mobile Number": "contact_data",
    "Voter ID":            "government_id",
    "Driving License":     "government_id",
}

# Document classification level → DLP decision mapping
DOC_CLASS_TO_DECISION = {
    RESTRICTED:   "BLOCK",
    CONFIDENTIAL: "BLOCK",
    INTERNAL:     "WARN",
    PUBLIC:       "PASS",
}

# Document classification level → risk score boost
DOC_CLASS_SCORE_MAP = {
    RESTRICTED:   99.0,
    CONFIDENTIAL: 85.0,
    INTERNAL:     40.0,
    PUBLIC:       5.0,
}


def check_dpdp(detected_types: List[str]):
    violations = [DPDP_SENSITIVE_MAP[t] for t in detected_types if t in DPDP_SENSITIVE_MAP]
    return len(violations) > 0, list(set(violations))


# ── Helpers ───────────────────────────────────────────────────────────────────
async def _get_prev_hash(session: AsyncSession) -> str:
    row = await session.scalar(
        select(DLPEvent.chain_hash).order_by(DLPEvent.id.desc()).limit(1)
    )
    return row or "GENESIS"


async def _update_user_profile(
    session, user_id, department, role, decision, risk_score
) -> None:
    profile = await session.scalar(
        select(UserRiskProfile).where(UserRiskProfile.user_id == user_id)
    )
    if not profile:
        profile = UserRiskProfile(
            user_id=user_id, department=department, role=role,
            total_prompts=0, total_blocked=0, total_warned=0, avg_risk_score=0.0,
        )
        session.add(profile)
    profile.total_prompts += 1
    profile.last_seen = datetime.utcnow()
    if decision == "BLOCK":
        profile.total_blocked += 1
    elif decision == "WARN":
        profile.total_warned += 1
    n = profile.total_prompts
    profile.avg_risk_score = (profile.avg_risk_score * (n - 1) + risk_score) / n


# ── Main prompt scan endpoint ─────────────────────────────────────────────────
@app.post("/gateway/analyze", response_model=AnalyzeResponse)
async def analyze(req: AnalyzeRequest, session: AsyncSession = Depends(get_session)):
    event_id  = str(uuid.uuid4())
    timestamp = datetime.utcnow().isoformat()

    cached = await get_cached(req.prompt)
    if cached:
        cached["event_id"]   = event_id
        cached["from_cache"] = True
        cached["timestamp"]  = timestamp
        return AnalyzeResponse(**cached)

    scan = await dlp_engine.scan(req.prompt, department=req.department)
    policy = policy_engine.evaluate(
        department=req.department,
        role=req.role,
        detected_types=scan.detected_types,
        risk_score=scan.risk_score,
    )

    final_decision = (
        "BLOCK" if policy["decision"] == "BLOCK" or scan.decision == "BLOCK"
        else "WARN" if policy["decision"] == "WARN" or scan.decision == "WARN"
        else "PASS"
    )

    dpdp_violation, dpdp_categories = check_dpdp(scan.detected_types)

    prev_hash  = await _get_prev_hash(session)
    chain_hash = DLPEvent.compute_chain_hash(prev_hash, event_id, timestamp, final_decision)

    db_event = DLPEvent(
        event_id=event_id,
        user_id=req.user_id,
        department=req.department,
        role=req.role,
        destination_model=req.destination_model,
        prompt_hash=hashlib.sha256(req.prompt.encode()).hexdigest(),
        prompt_snippet=req.prompt[:200],
        decision=final_decision,
        risk_score=scan.risk_score,
        risk_tier=scan.risk_tier,
        detected_types=scan.detected_types,
        block_reason=scan.block_reason,
        layer_scores=scan.layer_scores,
        findings_count=len(scan.findings),
        processing_ms=scan.processing_ms,
        from_cache=False,
        chain_hash=chain_hash,
        dpdp_violation=dpdp_violation,
        llm_triggered=scan.llm_triggered,
    )
    session.add(db_event)
    await _update_user_profile(
        session, req.user_id, req.department, req.role, final_decision, scan.risk_score
    )
    await session.commit()

    await alert_manager.evaluate(
        event_id=event_id,
        user_id=req.user_id,
        decision=final_decision,
        risk_score=scan.risk_score,
        detected_types=scan.detected_types,
        dpdp_violation=dpdp_violation,
    )

    result_dict = dict(
        event_id=event_id,
        decision=final_decision,
        risk_score=scan.risk_score,
        risk_tier=scan.risk_tier,
        block_reason=scan.block_reason,
        detected_types=scan.detected_types,
        findings=scan.findings,
        layer_scores=scan.layer_scores,
        processing_ms=scan.processing_ms,
        policy_notes=policy,
        from_cache=False,
        timestamp=timestamp,
        dpdp_violation=dpdp_violation,
        dpdp_categories=dpdp_categories,
        llm_triggered=scan.llm_triggered,
    )
    await set_cached(req.prompt, result_dict)
    return AnalyzeResponse(**result_dict)


# ── Document / File Upload Scan ───────────────────────────────────────────────
@app.post("/gateway/scan-file")
async def scan_file(
    user_id:     str = Form(...),
    department:  str = Form("default"),
    role:        str = Form("employee"),
    destination: str = Form("unknown"),
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_session),
):
    MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
    content  = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File too large. Max 10MB.")

    filename = file.filename or "unknown"
    ext      = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    text     = ""

    # ── Extract text ──────────────────────────────────────────────────────────
    try:
        if ext in ("txt", "csv", "json", "md", "py", "js", "env",
                   "yaml", "yml", "xml", "html", "log", "sql"):
            text = content.decode("utf-8", errors="ignore")

        elif ext == "pdf":
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(content)) as pdf:
                    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            except ImportError:
                raise HTTPException(status_code=422,
                    detail="pdfplumber not installed. Run: pip install pdfplumber")

        elif ext == "docx":
            try:
                from docx import Document
                doc  = Document(io.BytesIO(content))
                text = "\n".join(p.text for p in doc.paragraphs)
            except ImportError:
                raise HTTPException(status_code=422,
                    detail="python-docx not installed. Run: pip install python-docx")

        elif ext in ("xls", "xlsx"):
            try:
                import openpyxl
                wb   = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
                rows = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        rows.append(" ".join(str(c) for c in row if c is not None))
                text = "\n".join(rows)
            except ImportError:
                text = content.decode("utf-8", errors="ignore")
        else:
            text = content.decode("utf-8", errors="ignore")

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not parse file: {exc}")

    if not text.strip():
        return {
            "filename": filename,
            "decision": "PASS",
            "message":  "No readable text found in document.",
            "risk_score": 0,
            "doc_classification": "UNKNOWN",
            "doc_classification_confidence": 0.0,
        }

    # ── Document classification ───────────────────────────────────────────────
    # Determine if destination is an LLM
    is_llm_dest = any(d in destination.lower() for d in [
        "chatgpt", "gemini", "claude", "deepseek", "perplexity",
        "mistral", "copilot", "openai", "anthropic", "grok"
    ])

    doc_context = {
        "destination":   destination,
        "user_role":     role,
        "source_system": "",
        "off_hours":     False,
    }
    # Run both in parallel — LLM classification + DLP scan at the same time
    import asyncio as _asyncio
    classification, scan = await _asyncio.gather(
        doc_classifier.classify(text, filename=filename, context=doc_context),
        dlp_engine.scan(text, department=department),
    )
    dpdp_violation, dpdp_categories = check_dpdp(scan.detected_types)

    # ── Merge decisions ───────────────────────────────────────────────────────
    # Use the stricter of: content scan decision vs. document classification decision
    content_decision = scan.decision
    doc_decision     = DOC_CLASS_TO_DECISION.get(classification.level, "PASS")
    doc_risk_score   = DOC_CLASS_SCORE_MAP.get(classification.level, 0.0)

    # Take max risk
    final_risk_score = max(scan.risk_score, doc_risk_score)
    decision_rank    = {"BLOCK": 2, "WARN": 1, "PASS": 0}
    final_decision   = max(
        [content_decision, doc_decision],
        key=lambda d: decision_rank.get(d, 0)
    )

    # Build block reason
    block_reason = ""
    if final_decision == "BLOCK":
        if classification.level in (RESTRICTED, CONFIDENTIAL):
            block_reason = (
                f"Document classified as {classification.level} "
                f"({classification.confidence:.0%} confidence). "
                f"{classification.reasons[0] if classification.reasons else ''}"
            )
        else:
            block_reason = scan.block_reason or "Sensitive data detected in document."

    # ── Persist event ─────────────────────────────────────────────────────────
    event_id  = str(uuid.uuid4())
    timestamp = datetime.utcnow().isoformat()
    prev_hash = await _get_prev_hash(session)
    chain_hash = DLPEvent.compute_chain_hash(prev_hash, event_id, timestamp, final_decision)

    db_event = DLPEvent(
        event_id=event_id,
        user_id=user_id,
        department=department,
        role=role,
        destination_model=destination,
        prompt_hash=hashlib.sha256(content).hexdigest(),
        prompt_snippet=f"[FILE:{filename}] {text[:150]}",
        decision=final_decision,
        risk_score=final_risk_score,
        risk_tier=(
            "critical" if final_risk_score >= 80
            else "high"   if final_risk_score >= 60
            else "medium" if final_risk_score >= 30
            else "low"
        ),
        detected_types=scan.detected_types,
        block_reason=block_reason,
        layer_scores=scan.layer_scores,
        findings_count=len(scan.findings),
        processing_ms=scan.processing_ms,
        from_cache=False,
        chain_hash=chain_hash,
        dpdp_violation=dpdp_violation,
        doc_classification=classification.level,
        doc_classification_confidence=classification.confidence,
        doc_needs_review=classification.needs_review,
        doc_type=classification.matched_rules[0] if classification.matched_rules else None,
        llm_triggered=(scan.llm_triggered or classification.llm_triggered),
    )
    session.add(db_event)
    await _update_user_profile(session, user_id, department, role, final_decision, final_risk_score)
    await session.commit()

    # Alert for restricted docs
    if classification.level in (RESTRICTED, CONFIDENTIAL):
        await alert_manager.evaluate(
            event_id=event_id,
            user_id=user_id,
            decision=final_decision,
            risk_score=final_risk_score,
            detected_types=scan.detected_types,
            dpdp_violation=dpdp_violation,
            doc_classification=classification.level,
        )

    return {
        "filename":                      filename,
        "decision":                      final_decision,
        "block_reason":                  block_reason,
        "risk_score":                    final_risk_score,
        "risk_tier":                     db_event.risk_tier,
        "detected_types":                scan.detected_types,
        "findings":                      scan.findings[:10],
        "dpdp_violation":                dpdp_violation,
        "dpdp_categories":               dpdp_categories,
        "processing_ms":                 scan.processing_ms,
        # Document classification fields
        "doc_classification":            classification.level,
        "doc_classification_confidence": round(classification.confidence, 3),
        "doc_needs_review":              classification.needs_review,
        "doc_type":                      classification.doc_type,
        "doc_classification_method":     classification.method,
        "doc_classification_color":      classification.color,
        # Rich details for browser extension overlay
        "reasons":                       classification.reasons,
        "matched_rules":                 classification.matched_rules[:6],
        "pii_findings":                  classification.pii_findings,
        "llm_triggered":                 (scan.llm_triggered or classification.llm_triggered),
    }


# ── Admin endpoints (JWT protected) ──────────────────────────────────────────
@app.get("/admin/events")
async def list_events(
    limit:    int = Query(50, le=200),
    offset:   int = Query(0),
    decision: Optional[str] = None,
    session:  AsyncSession = Depends(get_session),
    _user:    dict = Depends(get_current_user),
):
    q = select(DLPEvent).order_by(DLPEvent.id.desc()).offset(offset).limit(limit)
    if decision:
        q = q.where(DLPEvent.decision == decision.upper())
    rows = (await session.scalars(q)).all()
    return [
        {
            "event_id":             r.event_id,
            "user_id":              r.user_id,
            "department":           r.department,
            "destination":          r.destination_model,
            "decision":             r.decision,
            "risk_score":           r.risk_score,
            "risk_tier":            r.risk_tier,
            "detected_types":       r.detected_types,
            "block_reason":         r.block_reason,
            "processing_ms":        r.processing_ms,
            "from_cache":           r.from_cache,
            "chain_hash":           r.chain_hash,
            "dpdp_violation":       r.dpdp_violation,
            "doc_classification":   r.doc_classification,
            "doc_needs_review":     r.doc_needs_review,
            "timestamp":            r.timestamp.isoformat() if r.timestamp else None,
        }
        for r in rows
    ]


@app.get("/admin/alerts")
async def list_alerts(
    limit:     int  = Query(50, le=200),
    dismissed: bool = False,
    session:   AsyncSession = Depends(get_session),
    _user:     dict = Depends(get_current_user),
):
    q = (
        select(Alert)
        .where(Alert.dismissed == dismissed)
        .order_by(Alert.id.desc())
        .limit(limit)
    )
    rows = (await session.scalars(q)).all()
    return [
        {
            "alert_id":          r.alert_id,
            "event_id":          r.event_id,
            "user_id":           r.user_id,
            "alert_type":        r.alert_type,
            "risk_score":        r.risk_score,
            "message":           r.message,
            "dismissed":         r.dismissed,
            "dpdp_violation":    r.dpdp_violation,
            "doc_classification": r.doc_classification,
            "timestamp":         r.timestamp.isoformat() if r.timestamp else None,
        }
        for r in rows
    ]


@app.post("/admin/alerts/{alert_id}/dismiss")
async def dismiss_alert(
    alert_id: str,
    session:  AsyncSession = Depends(get_session),
    _user:    dict = Depends(require_admin),
):
    alert = await session.scalar(select(Alert).where(Alert.alert_id == alert_id))
    if not alert:
        raise HTTPException(status_code=404, detail="Alert not found")
    alert.dismissed = True
    await session.commit()
    return {"status": "dismissed"}


@app.get("/admin/users")
async def list_users(
    limit:   int = Query(50, le=200),
    session: AsyncSession = Depends(get_session),
    _user:   dict = Depends(get_current_user),
):
    rows = (await session.scalars(
        select(UserRiskProfile).order_by(UserRiskProfile.total_blocked.desc()).limit(limit)
    )).all()
    return [
        {
            "user_id":       r.user_id,
            "department":    r.department,
            "role":          r.role,
            "total_prompts": r.total_prompts,
            "total_blocked": r.total_blocked,
            "total_warned":  r.total_warned,
            "avg_risk_score": round(r.avg_risk_score, 1),
            "last_seen":     r.last_seen.isoformat() if r.last_seen else None,
        }
        for r in rows
    ]


@app.get("/admin/users/{user_id}")
async def get_user(
    user_id: str,
    session: AsyncSession = Depends(get_session),
    _user:   dict = Depends(get_current_user),
):
    profile = await session.scalar(
        select(UserRiskProfile).where(UserRiskProfile.user_id == user_id)
    )
    if not profile:
        raise HTTPException(status_code=404, detail="User not found")
    recent_events = (await session.scalars(
        select(DLPEvent).where(DLPEvent.user_id == user_id)
        .order_by(DLPEvent.id.desc()).limit(20)
    )).all()
    return {
        "profile": {
            "user_id":       profile.user_id,
            "department":    profile.department,
            "role":          profile.role,
            "total_prompts": profile.total_prompts,
            "total_blocked": profile.total_blocked,
            "total_warned":  profile.total_warned,
            "avg_risk_score": round(profile.avg_risk_score, 1),
            "last_seen":     profile.last_seen.isoformat() if profile.last_seen else None,
        },
        "recent_events": [
            {
                "event_id":           e.event_id,
                "decision":           e.decision,
                "risk_score":         e.risk_score,
                "detected_types":     e.detected_types,
                "doc_classification": e.doc_classification,
                "timestamp":          e.timestamp.isoformat() if e.timestamp else None,
            }
            for e in recent_events
        ],
    }


@app.get("/admin/stats")
async def stats(
    session: AsyncSession = Depends(get_session),
    _user:   dict = Depends(get_current_user),
):
    total    = await session.scalar(select(func.count(DLPEvent.id)))
    blocked  = await session.scalar(select(func.count(DLPEvent.id)).where(DLPEvent.decision == "BLOCK"))
    warned   = await session.scalar(select(func.count(DLPEvent.id)).where(DLPEvent.decision == "WARN"))
    passed   = await session.scalar(select(func.count(DLPEvent.id)).where(DLPEvent.decision == "PASS"))
    alerts_n = await session.scalar(select(func.count(Alert.id)).where(Alert.dismissed == False))
    dpdp_n   = await session.scalar(select(func.count(DLPEvent.id)).where(DLPEvent.dpdp_violation == True))
    restricted_n = await session.scalar(
        select(func.count(DLPEvent.id)).where(DLPEvent.doc_classification == "RESTRICTED")
    )
    return {
        "total_prompts":    total    or 0,
        "total_blocked":    blocked  or 0,
        "total_warned":     warned   or 0,
        "total_passed":     passed   or 0,
        "active_alerts":    alerts_n or 0,
        "dpdp_violations":  dpdp_n   or 0,
        "restricted_docs":  restricted_n or 0,
        "block_rate_pct":   round((blocked or 0) / max(total or 1, 1) * 100, 1),
    }


@app.websocket("/ws/live")
async def ws_live(ws: WebSocket):
    await ws.accept()
    q = alert_manager.subscribe()
    try:
        while True:
            payload = await q.get()
            await ws.send_json(payload)
    except (WebSocketDisconnect, Exception):
        alert_manager.unsubscribe(q)


@app.get("/health")
async def health():
    return {"status": "ok", "service": "DLP Gateway", "version": "3.0.0"}
